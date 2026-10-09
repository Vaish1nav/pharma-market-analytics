import os
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE_DIR = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
PROCESSED_DATA_DIR = os.path.join(BASE_DIR, 'data', 'processed')
EXCEL_DIR = os.path.join(BASE_DIR, 'excel')
os.makedirs(EXCEL_DIR, exist_ok=True)

excel_path = os.path.join(EXCEL_DIR, 'pharma_market_analysis.xlsx')
print(f'Generating Executive Excel Report at {excel_path}...')

# Load processed CSVs
df_exec = pd.read_csv(os.path.join(PROCESSED_DATA_DIR, 'summary_executive_kpis.csv'))
df_prod = pd.read_csv(os.path.join(PROCESSED_DATA_DIR, 'summary_product_performance.csv'))
df_reg = pd.read_csv(os.path.join(PROCESSED_DATA_DIR, 'summary_regional_penetration.csv'))
df_price = pd.read_csv(os.path.join(PROCESSED_DATA_DIR, 'summary_pricing_index.csv'))
df_strat = pd.read_csv(os.path.join(PROCESSED_DATA_DIR, 'summary_strategic_matrix.csv'))

wb = openpyxl.Workbook()
# remove default sheet
wb.remove(wb.active)

# Styles
font_title = Font(name='Calibri', size=16, bold=True, color='FFFFFF')
font_subtitle = Font(name='Calibri', size=11, italic=True, color='595959')
font_header = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
font_bold = Font(name='Calibri', size=11, bold=True)
font_regular = Font(name='Calibri', size=11)
font_kpi_val = Font(name='Calibri', size=20, bold=True, color='1F497D')
font_kpi_lbl = Font(name='Calibri', size=9, bold=True, color='595959')

fill_title = PatternFill(start_color='1F497D', end_color='1F497D', fill_type='solid') # Corporate Navy
fill_header = PatternFill(start_color='2F5597', end_color='2F5597', fill_type='solid')
fill_kpi = PatternFill(start_color='F2F2F2', end_color='F2F2F2', fill_type='solid')
fill_accent = PatternFill(start_color='D9E1F2', end_color='D9E1F2', fill_type='solid')

border_thin = Side(border_style='thin', color='D9D9D9')
border_thick = Side(border_style='medium', color='1F497D')
box_border = Border(left=border_thin, right=border_thin, top=border_thin, bottom=border_thin)

# Function to style header
def style_table_header(ws, start_row, end_col):
    for col in range(1, end_col + 1):
        cell = ws.cell(row=start_row, column=col)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

# SHEET 1: Executive Summary
ws1 = wb.create_sheet(title='Executive Summary')
ws1.views.sheetView[0].showGridLines = True

ws1.merge_cells('A1:G1')
ws1['A1'] = "NovaPharm Laboratories - Portfolio & Market Performance Executive Summary"
ws1['A1'].font = font_title
ws1['A1'].fill = fill_title
ws1['A1'].alignment = Alignment(horizontal='left', vertical='center', indent=1)
ws1.row_dimensions[1].height = 40

# KPI Cards
kpis = [
    ("2025 TOTAL REVENUE", "$843.30 M", "B3:C4"),
    ("MARKET SHARE (2025)", "30.71 %", "D3:E4"),
    ("YoY GROWTH RATE", "+12.06 %", "F3:G4")
]

for title, val, rng in kpis:
    ws1.merge_cells(rng)
    top_cell = ws1[rng.split(':')[0]]
    top_cell.value = f"{title}\n\n{val}"
    top_cell.font = Font(name='Calibri', size=12, bold=True, color='1F497D')
    top_cell.fill = fill_kpi
    top_cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

# Executive Table
ws1.cell(row=6, column=1, value="Yearly Executive Summary").font = font_bold
headers1 = ["Year", "NovaPharm Revenue ($)", "NovaPharm Units", "Total Market Value ($)", "Value Market Share (%)"]
for col_num, h in enumerate(headers1, 1):
    ws1.cell(row=7, column=col_num, value=h)
style_table_header(ws1, 7, len(headers1))

for r_idx, row in df_exec.iterrows():
    r = 8 + r_idx
    ws1.cell(row=r, column=1, value=int(row['year'])).alignment = Alignment(horizontal='center')
    ws1.cell(row=r, column=2, value=row['novapharm_revenue']).number_format = '$#,##0.00'
    ws1.cell(row=r, column=3, value=row['novapharm_units']).number_format = '#,##0'
    ws1.cell(row=r, column=4, value=row['total_market_value']).number_format = '$#,##0.00'
    ws1.cell(row=r, column=5, value=row['value_market_share_pct'] / 100.0).number_format = '0.00%'

# Top 5 Products Table
ws1.cell(row=12, column=1, value="Top 5 Products by Revenue (2025)").font = font_bold
headers_top5 = ["Product Name", "Therapeutic Area", "Type", "2025 Revenue ($)", "YoY Growth (%)", "Market Share (%)"]
for col_num, h in enumerate(headers_top5, 1):
    ws1.cell(row=13, column=col_num, value=h)
style_table_header(ws1, 13, len(headers_top5))

for r_idx, row in df_prod.head(5).iterrows():
    r = 14 + r_idx
    ws1.cell(row=r, column=1, value=row['product_name'])
    ws1.cell(row=r, column=2, value=row['therapeutic_area'])
    ws1.cell(row=r, column=3, value=row['product_type'])
    ws1.cell(row=r, column=4, value=row['rev_2025']).number_format = '$#,##0.00'
    ws1.cell(row=r, column=5, value=row['yoy_growth_pct'] / 100.0).number_format = '0.00%'
    ws1.cell(row=r, column=6, value=row['market_share_2025_pct'] / 100.0).number_format = '0.00%'

# SHEET 2: Product Analysis
ws2 = wb.create_sheet(title='Product Analysis')
ws2.views.sheetView[0].showGridLines = True
ws2.merge_cells('A1:H1')
ws2['A1'] = "Product & Portfolio Performance Detail (2024 - 2025)"
ws2['A1'].font = font_title
ws2['A1'].fill = fill_title
ws2['A1'].alignment = Alignment(horizontal='left', vertical='center', indent=1)
ws2.row_dimensions[1].height = 35

headers2 = ["Product ID", "Product Name", "Brand Name", "Therapeutic Area", "Product Type", "2024 Revenue ($)", "2025 Revenue ($)", "YoY Growth (%)", "2025 Market Share (%)"]
for col_num, h in enumerate(headers2, 1):
    ws2.cell(row=3, column=col_num, value=h)
style_table_header(ws2, 3, len(headers2))

for r_idx, row in df_prod.iterrows():
    r = 4 + r_idx
    ws2.cell(row=r, column=1, value=row['product_id']).alignment = Alignment(horizontal='center')
    ws2.cell(row=r, column=2, value=row['product_name'])
    ws2.cell(row=r, column=3, value=row['brand_name'])
    ws2.cell(row=r, column=4, value=row['therapeutic_area'])
    ws2.cell(row=r, column=5, value=row['product_type'])
    ws2.cell(row=r, column=6, value=row['rev_2024']).number_format = '$#,##0.00'
    ws2.cell(row=r, column=7, value=row['rev_2025']).number_format = '$#,##0.00'
    ws2.cell(row=r, column=8, value=row['yoy_growth_pct'] / 100.0).number_format = '0.00%'
    ws2.cell(row=r, column=9, value=row['market_share_2025_pct'] / 100.0).number_format = '0.00%'

# SHEET 3: Regional Analysis
ws3 = wb.create_sheet(title='Regional Analysis')
ws3.views.sheetView[0].showGridLines = True
ws3.merge_cells('A1:G1')
ws3['A1'] = "Regional Market Performance & Penetration Analysis"
ws3['A1'].font = font_title
ws3['A1'].fill = fill_title
ws3['A1'].alignment = Alignment(horizontal='left', vertical='center', indent=1)
ws3.row_dimensions[1].height = 35

headers3 = ["Region ID", "Region", "Country", "Market Tier", "2025 Revenue ($)", "Addressable Population", "Patients Reached", "Penetration Rate (%)"]
for col_num, h in enumerate(headers3, 1):
    ws3.cell(row=3, column=col_num, value=h)
style_table_header(ws3, 3, len(headers3))

for r_idx, row in df_reg.iterrows():
    r = 4 + r_idx
    ws3.cell(row=r, column=1, value=row['region_id']).alignment = Alignment(horizontal='center')
    ws3.cell(row=r, column=2, value=row['region'])
    ws3.cell(row=r, column=3, value=row['country'])
    ws3.cell(row=r, column=4, value=row['market_type'])
    ws3.cell(row=r, column=5, value=row['region_revenue']).number_format = '$#,##0.00'
    ws3.cell(row=r, column=6, value=row['addressable_population']).number_format = '#,##0'
    ws3.cell(row=r, column=7, value=row['patients_reached']).number_format = '#,##0'
    ws3.cell(row=r, column=8, value=row['penetration_rate_pct'] / 100.0).number_format = '0.00%'

# SHEET 4: Pricing Analysis
ws4 = wb.create_sheet(title='Pricing Analysis')
ws4.views.sheetView[0].showGridLines = True
ws4.merge_cells('A1:H1')
ws4['A1'] = "Pricing Benchmarking & Price Index Analysis"
ws4['A1'].font = font_title
ws4['A1'].fill = fill_title
ws4['A1'].alignment = Alignment(horizontal='left', vertical='center', indent=1)
ws4.row_dimensions[1].height = 35

headers4 = ["Product ID", "Product Name", "Therapeutic Area", "Product Type", "NovaPharm Price ($)", "Competitor Price ($)", "Price Index", "Market Share (%)"]
for col_num, h in enumerate(headers4, 1):
    ws4.cell(row=3, column=col_num, value=h)
style_table_header(ws4, 3, len(headers4))

for r_idx, row in df_price.iterrows():
    r = 4 + r_idx
    ws4.cell(row=r, column=1, value=row['product_id']).alignment = Alignment(horizontal='center')
    ws4.cell(row=r, column=2, value=row['product_name'])
    ws4.cell(row=r, column=3, value=row['therapeutic_area'])
    ws4.cell(row=r, column=4, value=row['product_type'])
    ws4.cell(row=r, column=5, value=row['novapharm_price']).number_format = '$#,##0.00'
    ws4.cell(row=r, column=6, value=row['competitor_price']).number_format = '$#,##0.00'
    ws4.cell(row=r, column=7, value=row['price_index']).number_format = '0.0'
    ws4.cell(row=r, column=8, value=row['novapharm_market_share_pct'] / 100.0).number_format = '0.00%'

# SHEET 5: Strategic Opportunities
ws5 = wb.create_sheet(title='Strategic Opportunities')
ws5.views.sheetView[0].showGridLines = True
ws5.merge_cells('A1:H1')
ws5['A1'] = "Strategic Portfolio Allocation & Management Action Matrix"
ws5['A1'].font = font_title
ws5['A1'].fill = fill_title
ws5['A1'].alignment = Alignment(horizontal='left', vertical='center', indent=1)
ws5.row_dimensions[1].height = 35

headers5 = ["Product ID", "Product Name", "Therapeutic Area", "Revenue ($)", "Market Share (%)", "YoY Growth (%)", "Price Index", "Strategic Action"]
for col_num, h in enumerate(headers5, 1):
    ws5.cell(row=3, column=col_num, value=h)
style_table_header(ws5, 3, len(headers5))

for r_idx, row in df_strat.iterrows():
    r = 4 + r_idx
    ws5.cell(row=r, column=1, value=row['product_id']).alignment = Alignment(horizontal='center')
    ws5.cell(row=r, column=2, value=row['product_name'])
    ws5.cell(row=r, column=3, value=row['therapeutic_area'])
    ws5.cell(row=r, column=4, value=row['revenue']).number_format = '$#,##0.00'
    ws5.cell(row=r, column=5, value=row['market_share_pct'] / 100.0).number_format = '0.00%'
    ws5.cell(row=r, column=6, value=row['yoy_growth_pct'] / 100.0).number_format = '0.00%'
    ws5.cell(row=r, column=7, value=row['price_index']).number_format = '0.0'
    
    act_cell = ws5.cell(row=r, column=8, value=row['strategic_recommendation'])
    act_cell.font = font_bold
    if 'INVEST' in str(row['strategic_recommendation']):
        act_cell.fill = PatternFill(start_color='C6EFCE', end_color='C6EFCE', fill_type='solid') # Soft green
    elif 'OPTIMIZE' in str(row['strategic_recommendation']):
        act_cell.fill = PatternFill(start_color='FFEB9C', end_color='FFEB9C', fill_type='solid') # Soft yellow
    elif 'MONITOR' in str(row['strategic_recommendation']):
        act_cell.fill = PatternFill(start_color='FFC7CE', end_color='FFC7CE', fill_type='solid') # Soft red

# Auto-adjust column widths across all sheets
for sheet in wb.worksheets:
    for col in sheet.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.row == 1:
                continue
            if cell.value:
                val_str = str(cell.value)
                if len(val_str) > max_len:
                    max_len = len(val_str)
        sheet.column_dimensions[col_letter].width = max(max_len + 4, 12)

wb.save(excel_path)
print(f'Workbook saved successfully with 5 formatted sheets!')
